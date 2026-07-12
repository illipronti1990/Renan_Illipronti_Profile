#!/usr/bin/env python3
"""Converte planilha .xlsx (openpyxl) para formato .xls (Excel 97-2003)."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

import xlwt
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from xlwt import Formula


def _excel_color(hex_rgb: str | None) -> int | None:
    if not hex_rgb:
        return None
    rgb = hex_rgb.strip().lstrip("#")
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return None
    try:
        int(rgb[0:2], 16)
        int(rgb[2:4], 16)
        int(rgb[4:6], 16)
    except ValueError:
        return None
    palette = {
        "1B2A4A": 55,
        "C9A227": 51,
        "FFFACD": 43,
        "E8EEF4": 22,
        "F5F5F5": 22,
        "FFFFFF": 1,
        "FFC7CE": 10,
        "C6EFCE": 11,
        "FFEB9C": 43,
        "FCE4D6": 53,
        "2E7D32": 17,
        "C62828": 10,
        "000000": 0,
    }
    return palette.get(rgb.upper(), 22)


def _build_style(cell, book: xlwt.Workbook, cache: dict) -> xlwt.XFStyle:
    key = (
        cell.font.bold if cell.font else False,
        cell.font.size if cell.font else 11,
        cell.font.color.rgb if cell.font and cell.font.color and cell.font.color.rgb else None,
        cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb else None,
        cell.number_format,
        cell.alignment.horizontal if cell.alignment else None,
    )
    if key in cache:
        return cache[key]

    style = xlwt.XFStyle()
    font = xlwt.Font()
    font.name = "Calibri"
    font.bold = bool(cell.font and cell.font.bold)
    font.height = int((cell.font.size if cell.font and cell.font.size else 11) * 20)
    if cell.font and cell.font.color and cell.font.color.rgb:
        rgb = str(cell.font.color.rgb)
        color_idx = _excel_color(rgb[-6:] if len(rgb) >= 6 else rgb)
        if color_idx is not None:
            font.colour_index = color_idx
    style.font = font

    if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
        rgb = str(cell.fill.fgColor.rgb)
        if len(rgb) >= 6:
            bg = _excel_color(rgb[-6:])
            if bg:
                pattern = xlwt.Pattern()
                pattern.pattern = xlwt.Pattern.SOLID_PATTERN
                pattern.pattern_fore_colour = bg
                style.pattern = pattern

    align = xlwt.Alignment()
    horiz = cell.alignment.horizontal if cell.alignment else None
    if horiz == "center":
        align.horz = xlwt.Alignment.HORZ_CENTER
    elif horiz == "right":
        align.horz = xlwt.Alignment.HORZ_RIGHT
    else:
        align.horz = xlwt.Alignment.HORZ_LEFT
    align.wrap = 1
    style.alignment = align

    nf = cell.number_format or "General"
    if "R$" in nf or "#,##0.00" in nf:
        style.num_format_str = "#,##0.00"
    elif nf in ("DD/MM/YYYY", "dd/mm/yyyy"):
        style.num_format_str = "DD/MM/YYYY"
    elif "0%" in nf or "0.0%" in nf:
        style.num_format_str = nf.replace("0.0%", "0.00%")

    cache[key] = style
    return style


FORMULA_REPLACEMENTS = [
    ("SOMASES", "SUMIFS"),
    ("SOMASE", "SUMIF"),
    ("SEERRO", "IFERROR"),
    ("FIM.MÊS", "EOMONTH"),
    ("PROCV", "VLOOKUP"),
    ("ÍNDICE", "INDEX"),
    ("CORRESP", "MATCH"),
    ("MAIOR", "LARGE"),
    ("MÊS", "MONTH"),
    ("HOJE", "TODAY"),
    ("FALSO", "FALSE"),
    ("VERDADEIRO", "TRUE"),
    ("SOMA", "SUM"),
    ("DATA", "DATE"),
    ("TEXTO", "TEXT"),
    ("FILTRAR", "FILTER"),
    ("LIN", "ROW"),
]


def _to_xls_formula(formula: str) -> str:
    result = formula.lstrip("=")
    for pt, en in FORMULA_REPLACEMENTS:
        result = result.replace(pt, en)
    result = re.sub(r"\bSE\b", "IF", result)
    result = result.replace(";", ",")
    # IFERROR → IF(ISERROR(...)) para compatibilidade Excel 97-2003
    result = re.sub(
        r"IFERROR\(([^,]+),([^)]+)\)",
        r"IF(ISERROR(\1),\2,\1)",
        result,
        flags=re.IGNORECASE,
    )
    return result


def _cell_write_value(cell) -> object:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        try:
            return Formula(_to_xls_formula(value))
        except Exception:
            return value
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return value


def xlsx_to_xls(xlsx_path: Path, xls_path: Path) -> Path:
    src = load_workbook(xlsx_path, data_only=False)
    book = xlwt.Workbook(encoding="utf-8")
    style_cache: dict = {}

    # Passo 1: criar todas as abas (necessário para referências cruzadas)
    out_sheets: dict[str, xlwt.Worksheet] = {}
    for sheet_name in src.sheetnames:
        out = book.add_sheet(sheet_name[:31])
        if src[sheet_name].sheet_state == "hidden":
            out.visibility = 1
        out_sheets[sheet_name] = out

    # Passo 2: preencher células
    for sheet_name in src.sheetnames:
        ws = src[sheet_name]
        out = out_sheets[sheet_name]

        merge_top_left: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for merged in ws.merged_cells.ranges:
            top_left = (merged.min_row, merged.min_col)
            bounds = (merged.min_row, merged.min_col, merged.max_row, merged.max_col)
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    merge_top_left[(r, c)] = bounds

        written_merges: set[tuple[int, int, int, int]] = set()

        for row in ws.iter_rows():
            for cell in row:
                r, c = cell.row, cell.column
                style = _build_style(cell, book, style_cache)
                value = _cell_write_value(cell)

                if (r, c) in merge_top_left:
                    r1, c1, r2, c2 = merge_top_left[(r, c)]
                    bounds = (r1, c1, r2, c2)
                    if (r, c) != (r1, c1):
                        continue
                    if bounds in written_merges:
                        continue
                    written_merges.add(bounds)
                    out.write_merge(r1 - 1, r2 - 1, c1 - 1, c2 - 1, value, style)
                else:
                    out.write(r - 1, c - 1, value, style)

        for col_letter, dim in ws.column_dimensions.items():
            if dim.width:
                try:
                    idx = column_index_from_string(col_letter) - 1
                    out.col(idx).width = int(dim.width * 256)
                except ValueError:
                    pass

        for row_idx, dim in ws.row_dimensions.items():
            if dim.height:
                out.row(row_idx - 1).height = int(dim.height * 20)

    book.save(str(xls_path))
    return xls_path


if __name__ == "__main__":
    base = Path(__file__).parent
    src = base / "Tesouraria_Igreja.xlsx"
    dst = base / "Tesouraria_Igreja.xls"
    if not src.exists():
        from gerar_planilha import generate

        generate(src)
    path = xlsx_to_xls(src, dst)
    print(f"Arquivo .xls gerado: {path}")
